"""OASIS proje özeti raporunu DOCX olarak üretir."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

import os

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "OASIS_COP31_Proje_Ozeti.docx")

ACCENT = RGBColor(0x0F, 0x4C, 0x81)
GRAY = RGBColor(0x44, 0x44, 0x44)

doc = Document()

# Sayfa düzeni
for s in doc.sections:
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.4)
    s.right_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12


def heading(text, size=13, space_before=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = ACCENT
    return p


def body(text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(val)
            r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ─────────────────────────── BAŞLIK ───────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("OASIS")
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = ACCENT
title.paragraph_format.space_after = Pt(0)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Open Awareness System for Incident Support")
r.font.size = Pt(11.5)
r.italic = True
r.font.color.rgb = GRAY
sub.paragraph_format.space_after = Pt(0)

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run("Yapay Zekâ Destekli Afet Yönetim Karar Destek Sistemi")
r.bold = True
r.font.size = Pt(12)
sub2.paragraph_format.space_after = Pt(2)

ctx = doc.add_paragraph()
ctx.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ctx.add_run(
    "COP31'e Doğru Hacettepe Üniversitesi İklim Diyalogları\n"
    "Dirençli Kentler, Doğa Temelli Çözümler ve Dijital Risk Yönetimi"
)
r.font.size = Pt(9.5)
r.font.color.rgb = GRAY
ctx.paragraph_format.space_after = Pt(10)

# ─────────────────────────── ÖZET ───────────────────────────
heading("1. Yönetici Özeti", space_before=0)
body(
    "OASIS, afet anında sahadan gelen vatandaş ihbarlarını saniyeler içinde standart, ölçülebilir ve "
    "haritalanabilir operasyonel rapora dönüştüren uçtan uca bir karar destek sistemidir. Sistem, "
    "vatandaşın telefonundan gönderdiği kısa metni, fotoğrafı ve konumu alır. Bu ham veriyi yapılandırılmış "
    "bir olay kaydına çevirir. Aciliyet puanını deterministik bir formülle hesaplar. Sonucu gerçek bina "
    "modelleri üzerinde üç boyutlu bir kent sahnesinde karar vericiye sunar. Amaç, afetin ilk saatlerinde "
    "kaybedilen en değerli kaynağı korumaktır: zaman."
)
body(
    "Sistemin ayırt edici yanı, yapay zekâyı karar merciine değil, rapor düzeltme ve standartlaştırma "
    "katmanına yerleştirmesidir. Kararın kendisi şeffaf ve tekrarlanabilir kurallarla üretilir. Böylece "
    "karar vericiye tahmin değil, doğrulanabilir bir rapor sunulur."
)

# ─────────────────────────── PROBLEM ───────────────────────────
heading("2. Çözülen Problem")
body(
    "İklim kaynaklı afetlerde sel, heyelan, fırtına ve aşırı yağış olayları kentsel alanlarda giderek "
    "sıklaşıyor. Bu olaylarda kritik darboğaz genellikle müdahale kapasitesi değil, durumsal farkındalık "
    "eksikliğidir. İlk saatlerde afet yönetim merkezine ulaşan bilgi dağınıktır. Telefon hatları kilitlenir. "
    "Sosyal medya doğrulanamaz. Sahadan gelen bilgi serbest metindir ve karşılaştırılabilir değildir. "
    "Sonuç olarak ekipler nereye önce gideceğine eksik bilgiyle karar verir."
)
body(
    "OASIS bu boşluğu kapatır. Vatandaşı pasif bir yardım talep edeni olmaktan çıkarıp kalibre edilmiş bir "
    "veri kaynağına dönüştürür. Yüzlerce dağınık ihbarı, önceliklendirilmiş olay kümelerine indirger."
)

# ─────────────────────────── NASIL ÇALIŞIR ───────────────────────────
heading("3. Sistem Nasıl Çalışır")
bullet(
    "Vatandaş mobil uygulamadan mağdur veya gözlemci olarak ihbar oluşturur. Kısa bir metin girer, "
    "isterse fotoğraf çeker. Konum otomatik eklenir.",
    bold_lead="İhbar. ",
)
bullet(
    "Metin ve görsel birlikte değerlendirilir. Sistem 18 adet standart afet göstergesini çıkarır. "
    "Mahsur kişi var mı, bina çökmüş mü, yol kapalı mı, su yükseliyor mu gibi net sorular yanıtlanır.",
    bold_lead="Yapılandırma. ",
)
bullet(
    "Göstergeler ağırlıklı ve sınırlandırılmış bir formülle 1 ile 10 arasında aciliyet puanına çevrilir. "
    "İnsan faktörü en yüksek ağırlığa sahiptir. Ayrıca verinin ne kadar güvenilir olduğunu belirten ayrı "
    "bir güvenilirlik skoru üretilir.",
    bold_lead="Puanlama. ",
)
bullet(
    "Birbirine yakın ve aynı türden ihbarlar mekânsal kümeleme ile tek bir olaya bağlanır. Böylece elli "
    "ihbar elli ayrı görev değil, beş gerçek olay olarak görünür. Her kümenin merkezi, yarıçapı ve "
    "ortalama şiddeti otomatik güncellenir.",
    bold_lead="Kümeleme. ",
)
bullet(
    "Sonuç, gerçek bina geometrileri üzerinde üç boyutlu bir kent sahnesinde gösterilir. Karar verici "
    "olayı listede değil, mekânda görür.",
    bold_lead="Sunum. ",
)
body(
    "Tüm bu zincir kullanıcıyı bekletmeden arka planda çalışır. İhbarı gönderen kişi için işlem anında "
    "biter. Karar verici için olay, gönderimden saniyeler sonra haritada belirir."
)

# ─────────────────────────── YAPAY ZEKÂNIN ROLÜ ───────────────────────────
heading("4. Yapay Zekânın Rolü: Karar Değil, Doğru Rapor")
body(
    "Afet yönetiminde yapay zekâya yöneltilen en haklı eleştiri, kararın izlenemez hale gelmesidir. OASIS "
    "bu riski tasarım düzeyinde ortadan kaldırır. Yapay zekâ modeli hangi ekibin nereye gideceğine karar "
    "vermez. Modelin tek görevi, panik içinde yazılmış, imla hatalarıyla dolu, eksik ve düzensiz metni "
    "standart bir olay kaydına çevirmektir. Yani yapay zekâ burada bir karar vericinin değil, bir raporlama "
    "editörünün işini yapar."
)
body(
    "Aciliyet puanı ve önerilen aksiyon, model çıktısından değil, sabit ağırlıklara sahip deterministik bir "
    "kural setinden üretilir. Aynı göstergeler her zaman aynı puanı verir. Karar verici puanın nasıl "
    "oluştuğunu satır satır geriye doğru izleyebilir. Bu, sistemin resmi bir müdahale sürecinde "
    "kullanılabilmesi için zorunlu olan hesap verebilirliği sağlar."
)
body(
    "Rapor doğruluğu bağımsız testlerle ölçülmüştür. Bilinen cevap testinde bina çökmesi göstergesi tam "
    "isabetle, su yükselmesi ve hayati tehlike göstergeleri yüksek başarımla tespit edilmiştir. Farklı "
    "modellerin aynı ihbar üzerindeki uzlaşı oranı çoğu vakada yüzde 80 seviyesinin üzerindedir. Metin "
    "bozma testlerinde büyük harf, Türkçe karakter kaybı ve yazım hatası eklendiğinde sistemin ürettiği "
    "gösterge kümesi büyük ölçüde sabit kalmıştır. Bu, gerçek afet metinlerinin bozuk yazıldığı düşünülürse "
    "sistemin en kritik dayanıklılık özelliğidir."
)

# ─────────────────────────── ÜÇ BOYUT ───────────────────────────
heading("5. Üç Boyutlu Kent Modeli ve Gerçek Bina Geometrileri")
body(
    "OASIS'in karar destek arayüzü iki boyutlu bir nokta haritası değildir. Sistem, resmi kaynaklardan "
    "sağlanan gerçek bina modellerini, sayısal arazi modelini ve yüksek çözünürlüklü ortofotoyu birlikte "
    "kullanır. Kent, olduğu gibi üç boyutlu olarak ekrana gelir."
)
bullet(
    "İhbar bir mahalleye değil, belirli bir binaya bağlanır. Karar verici hangi yapının kaç katlı "
    "olduğunu, çevresindeki yolların durumunu ve arazinin eğimini doğrudan görür.",
    bold_lead="Bina ölçeğinde çözünürlük. ",
)
bullet(
    "Binalar altı kademeli bir hasar ve etkilenme skalasına göre renklendirilir. Ağır etkilenen yapılar "
    "sahnede anında ayırt edilir.",
    bold_lead="Hasar seviyesi görselleştirmesi. ",
)
bullet(
    "Su seviyesi kademeli olarak yükseltilerek hangi yapıların ve yolların hangi kotta etkileneceği "
    "önceden görülebilir. Bu, tahliye rotası ve toplanma alanı seçimini somut hale getirir.",
    bold_lead="Sel simülasyonu. ",
)
bullet(
    "Zaman filtresi ile son bir saatte, son altı saatte veya seçilen tarih aralığında ne olduğu "
    "izlenebilir. Olayın mekânda nasıl büyüdüğü görülür.",
    bold_lead="Zaman boyutu. ",
)
body(
    "Bu üç boyutlu temsil bir görsellik tercihi değildir. Afet müdahalesinde kat sayısı, bina yüksekliği, "
    "arazi eğimi ve su kotu doğrudan operasyonel değişkenlerdir. İki boyutlu harita bu bilgiyi taşıyamaz."
)

# ─────────────────────────── İLETİŞİM ÇÖKÜŞÜ ───────────────────────────
heading("6. İletişim Altyapısı Çöktüğünde")
body(
    "Afetin ilk saatlerinde baz istasyonları aşırı yüklenir veya devre dışı kalır. Sesli arama ve mesajlaşma "
    "pratikte kullanılamaz hale gelir. OASIS bu senaryo düşünülerek tasarlanmıştır. Bir ihbar, çok küçük bir "
    "veri paketinden ibarettir. Kısa metin, konum ve sıkıştırılmış tek bir fotoğraf gönderilir. Görseller "
    "cihaz üzerinde küçültülerek iletilir. Bu sayede telefon ve mesajlaşma altyapısının çöktüğü, yalnızca çok "
    "zayıf bir internet bağlantısının kaldığı koşullarda dahi rapor gönderimi yapılabilir."
)
body(
    "Sistem ayrıca kullanıcının iletişim kurabilme durumunu ayrı bir alan olarak kaydeder. Yazamayacak "
    "durumdaki bir mağdurun kaydı ile bir gözlemcinin kaydı aynı biçimde değerlendirilmez."
)
body(
    "Gelecek aşamada Bluetooth tabanlı cihazdan cihaza iletim planlanmaktadır. Şebekesi olmayan bir "
    "kullanıcının ihbarı, yakınındaki cihazlar üzerinden atlayarak bağlantısı olan ilk cihaza ulaşacak ve "
    "oradan afet yönetim ekibine iletilecektir. Böylece afet alanındaki insanlar, hiçbir şebeke olmasa bile "
    "afet yönetim ekibine hızlıca ulaşabilecektir."
)

# ─────────────────────────── FARK ───────────────────────────
heading("7. Mevcut Sistemlerden Farkı")
table(
    ["Boyut", "Yaygın uygulama", "OASIS yaklaşımı"],
    [
        ["Girdi", "Çağrı merkezi, serbest metin, sosyal medya", "Yapılandırılmış ihbar, konum ve fotoğraf birlikte"],
        ["Önceliklendirme", "Operatörün yorumu, değişken", "Sabit ağırlıklı deterministik puan, tekrarlanabilir"],
        ["Yapay zekâ", "Karar öneren kara kutu", "Rapor standartlaştıran, izlenebilir katman"],
        ["Görselleştirme", "İki boyutlu nokta haritası", "Gerçek bina modelleriyle üç boyutlu kent sahnesi"],
        ["Yığın ihbar", "Her ihbar ayrı kayıt", "Mekânsal kümeleme ile tek olaya indirgeme"],
        ["Bant genişliği", "Sürekli bağlantı varsayımı", "Çok düşük veri ile çalışır, şebeke dışı iletim hedefi"],
    ],
    widths=[3.0, 5.4, 6.6],
)

# ─────────────────────────── FAYDA ───────────────────────────
heading("8. Dünyaya Katkısı ve İklim Dirençliliği Bağlamı")
body(
    "OASIS'in katkısı tek bir kuruma veya tek bir afet türüne bağlı değildir. Sistem, iklim kaynaklı "
    "afetlerde tekrarlayan üç evrensel probleme yanıt üretir."
)
bullet(
    "Sahadan gelen dağınık ve doğrulanamayan bilgi, dakikalar yerine saniyeler içinde "
    "önceliklendirilmiş operasyonel tabloya dönüşür. Kurtarma ekipleri kaynağı doğru yere yönlendirir.",
    bold_lead="Müdahale hızı. ",
)
bullet(
    "Puanlama kuralları açıktır. Her karar geriye doğru izlenebilir. Bu, kamu kurumlarının "
    "hesap verebilirlik gereksinimini karşılar ve sistemin kurumsal olarak benimsenmesini mümkün kılar.",
    bold_lead="Şeffaflık. ",
)
bullet(
    "Biriken ihbar ve küme verisi, kentin hangi noktalarının hangi olayda tekrar tekrar etkilendiğini "
    "gösterir. Bu birikim, risk haritalarının güncellenmesine, doğa temelli çözümlerin doğru konuma "
    "yerleştirilmesine ve imar kararlarının veriye dayandırılmasına girdi sağlar. Afet anında toplanan veri, "
    "afet öncesi planlamayı besleyen kalıcı bir kaynağa dönüşür.",
    bold_lead="Planlamaya geri besleme. ",
)
body(
    "Coğrafi bilgi sistemleri ve mekânsal yapay zekânın iklim dirençliliğine katkısı tam da bu noktada "
    "somutlaşır. Mekânsal veri yalnızca olayı haritalamak için değil, müdahaleyi yönetmek ve sonrasında "
    "kenti yeniden planlamak için kullanılır. OASIS bu döngünün çalışan bir örneğidir."
)

# ─────────────────────────── DURUM ───────────────────────────
heading("9. Mevcut Durum ve Yol Haritası")
body(
    "Sistem çalışır durumda bir prototiptir. Mobil uygulama, sunucu tarafı, mekânsal veritabanı, analiz "
    "katmanı ve üç boyutlu karar destek arayüzü uçtan uca entegre edilmiştir. Doğrulama testleri gerçek "
    "senaryolara dayalı test veri seti üzerinde yürütülmüştür. Kısa vadede planlanan çalışmalar; şebeke dışı "
    "Bluetooth iletim, sanal gerçeklik ortamında afet sahnesi rekonstrüksiyonu ve kurumsal ekip yönetimi "
    "modülünün eklenmesidir."
)

# ─────────────────────────── SUNUM ÖZETİ ───────────────────────────
heading("10. Sunum İçin Kısa Özet")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.left_indent = Cm(0.5)
p.paragraph_format.right_indent = Cm(0.5)
r = p.add_run(
    "OASIS, iklim kaynaklı afetlerde vatandaştan gelen ihbarı saniyeler içinde standart bir operasyonel "
    "rapora çeviren bir karar destek sistemidir. Panik içinde yazılmış bir metni, fotoğrafı ve konumu alır; "
    "yapay zekâyı karar vermek için değil, raporu doğru ve karşılaştırılabilir hale getirmek için kullanır. "
    "Aciliyet puanı şeffaf ve tekrarlanabilir kurallarla hesaplanır, bu nedenle karar verici tahmin değil "
    "izlenebilir bir rapor görür. Birbirine yakın ihbarlar mekânsal olarak kümelenir ve sonuç, gerçek bina "
    "modelleriyle kurulmuş üç boyutlu bir kent sahnesinde bina ölçeğinde sunulur. Sistem çok düşük veri "
    "kullandığı için telefon ve mesajlaşma altyapısının çöktüğü koşullarda bile ihbar iletimine izin verir; "
    "gelecek aşamada Bluetooth tabanlı cihazdan cihaza iletim ile şebeke olmadan da afet yönetim ekibine "
    "ulaşım hedeflenmektedir. Toplanan veri afet sonrasında risk haritalarını ve kentsel planlama kararlarını "
    "besleyen kalıcı bir girdiye dönüşür."
)
r.italic = True
r.font.size = Pt(10)

foot = doc.add_paragraph()
foot.paragraph_format.space_before = Pt(14)
r = foot.add_run(
    "Hazırlayan: OASIS Proje Ekibi  |  Harita Mühendisliği Bitirme Projesi  |  Ağustos 2026"
)
r.font.size = Pt(8.5)
r.font.color.rgb = GRAY
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.save(OUT)
print("saved:", OUT)
